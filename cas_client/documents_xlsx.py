"""Excel export for the caja's "cartera por vencer" report
(GetUpcomingDueReport / documents-equivalent of GetClientPaymentStatusReport,
but looking forward instead of at what's already overdue).

Deliberately the only report in the app with **no** PDF/DOCX/print
counterpart: it was requested as "solamente un excel" because a teller
working a 7-day call list needs sortable/filterable cells (by date, by
amount) more than a formatted page -- so this builds an .xlsx directly with
openpyxl instead of going through documents.py's HTML-string pipeline the way
every other document does.

Amounts and dates are written as real numeric/date cell values (not display
strings like formatting.gs()/formatting.fecha()) so the sheet stays usable
inside Excel -- sorting by "Vencimiento" or summing "Monto (Gs)" has to work
without the user reformatting the column first.
"""

from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from cas_client import theme
from cas_client.formatting import fecha_hora

_HEADERS = (
    "Cliente",
    "Celular",
    "Préstamo",
    "N° cuota",
    "Vencimiento",
    "Monto de la cuota (Gs)",
    "Cuotas ya abonadas",
    "Plazo (cuotas)",
)

# Sin decimales: el guaraní no tiene subdivisión en la práctica, mismo
# criterio que formatting.gs().
_MONEY_FORMAT = '#,##0" Gs"'
_DATE_FORMAT = "dd/mm/yyyy"

_HEADER_FILL = PatternFill("solid", fgColor=theme.PRIMARY.lstrip("#"))
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=13, color=theme.PRIMARY.lstrip("#"))
_SUBTITLE_FONT = Font(italic=True, color=theme.TEXT_MUTED.lstrip("#"))


def reporte_cartera_por_vencer_workbook(response) -> Workbook:
    """Arma el libro a partir de una GetUpcomingDueReportResponse.

    Una fila por CUOTA dentro de la ventana, no por cliente ni por préstamo
    -- mismo criterio que la respuesta del servidor (ver
    dashboard_service.GetUpcomingDueReport): un cliente con dos préstamos
    activos, cada uno con una cuota próxima, aparece dos veces.
    """
    libro = Workbook()
    hoja: Worksheet = libro.active
    hoja.title = "Cartera por vencer"

    ultima_columna = get_column_letter(len(_HEADERS))

    hoja.merge_cells(f"A1:{ultima_columna}1")
    titulo = hoja["A1"]
    titulo.value = f"Cartera por vencer en los próximos {response.days_ahead} días"
    titulo.font = _TITLE_FONT

    hoja.merge_cells(f"A2:{ultima_columna}2")
    subtitulo = hoja["A2"]
    subtitulo.value = f"Generado el {fecha_hora(response.generated_at.ToDatetime())}"
    subtitulo.font = _SUBTITLE_FONT

    fila_encabezado = 4
    for columna, texto in enumerate(_HEADERS, start=1):
        celda = hoja.cell(row=fila_encabezado, column=columna, value=texto)
        celda.font = _HEADER_FONT
        celda.fill = _HEADER_FILL
        celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)

    fila = fila_encabezado
    for dato in response.rows:
        fila += 1
        anio, mes, dia = (int(parte) for parte in dato.due_date.split("-"))
        valores = (
            dato.client_name,
            dato.phone_number,
            dato.loan_id[:8],
            dato.installment_number,
            date(anio, mes, dia),
            Decimal(dato.due_amount),
            dato.installments_paid_count,
            dato.term_months,
        )
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)
            if columna == 5:  # Vencimiento
                celda.number_format = _DATE_FORMAT
            elif columna == 6:  # Monto de la cuota
                celda.number_format = _MONEY_FORMAT

    if not response.rows:
        hoja.cell(
            row=fila_encabezado + 1,
            column=1,
            value="No hay cuotas por vencer en esa ventana.",
        )

    for columna, encabezado in enumerate(_HEADERS, start=1):
        letra = get_column_letter(columna)
        ancho = max(len(encabezado) + 2, 14)
        hoja.column_dimensions[letra].width = ancho

    return libro
